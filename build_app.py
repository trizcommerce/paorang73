"""셀러 콘텐츠 캘린더 앱 빌더. (셀러별 설정은 config.json)

구글 시트(공개 xlsx export)를 내려받아 캘린더/제품/월별 데이터를 JSON으로 정리한 뒤
template.html 에 넣어 모바일용 단일 HTML(duckduck_calendar.html)을 만든다.

    python build_app.py            # 시트 새로 받아서 빌드
    python build_app.py --local    # 이미 받아둔 sheet.xlsx 로 빌드
"""
import datetime as dt
import json
import os
import re
import sys
import urllib.request
from pathlib import Path

import openpyxl

import sheet_media

HERE = Path(__file__).parent
CONFIG = json.loads((HERE / "config.json").read_text(encoding="utf-8"))
SHEET_ID = CONFIG["sheet_id"]
YEAR = CONFIG.get("year", 2026)
XLSX = HERE / "sheet.xlsx"
TEMPLATE = HERE / "template.html"
OUT = HERE / "duckduck_calendar.html"
MEDIA = HERE / "media"  # 시트 이미지를 줄여서 저장하는 곳 (앱에서는 media/파일명)
IMAGES = {}  # 시트 이름 → {(행, 열): [파일명]}
FORMATS = {}  # (시트 이름, 행, 열) → 서식·링크 정보 (sheet_media.CellFormat)

CAL_COLS = "BCDEFGH"  # 월~일
MEETING_FILL = "FFCFE2F3"
PALETTE = ["#A4C2F4", "#F9CB9C", "#B6D7A8", "#D9D2E9", "#FFE599", "#EA9999", "#A2C4C9", "#D5A6BD"]
HOLIDAY_FILL = "FFF3F3F3"


TABS_DIR = HERE / "tabs"  # 탭별로 받은 xlsx (config "per_tab": true)


def fetch(url, dest, tries=4):
    """파일 받기 — 중간에 끊기면 다시 시도, 덜 받은 xlsx 는 걸러냄"""
    import time
    import zipfile
    for i in range(1, tries + 1):
        try:
            with urllib.request.urlopen(url, timeout=600) as r:
                dest.write_bytes(r.read())
            zipfile.ZipFile(dest).testzip()
            return
        except Exception as e:
            print(f"download try {i} failed: {e}")
            if i == tries:
                raise
            time.sleep(20 * i)


def load_tabs(local):
    """탭별로 받아 [(워크북, 파일)] — 시트 순서대로 받다가 올해 이전 공구 탭이 3개 연속이면 멈춤
    (한 번에 받기엔 너무 큰 시트용)"""
    TABS_DIR.mkdir(exist_ok=True)
    out, old = [], 0
    by_prefix = CONFIG.get("tabs") == "year_prefix"
    for name, gid in sheet_media.list_tabs(SHEET_ID):
        if by_prefix:  # "26.08_…" 처럼 이름에 연월이 붙은 탭 중 올해 것만 (연월 없는 탭은 옛 탭)
            ym = TAB_YM_RE.match(name.strip("[] "))
            if not ym or 2000 + int(ym[1]) < YEAR:
                continue
        f = TABS_DIR / f"{gid}.xlsx"
        if not (local and f.exists()):
            fetch(f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx&gid={gid}", f)
        wb = openpyxl.load_workbook(f, data_only=True, rich_text=True)
        ws = wb.worksheets[0]
        if find_header(ws):
            p = parse_product(ws, {})
            dates = [s["date"] for s in p["slots"] if s["date"]] + (p["period"] or [])
            if dates and max(dates) < f"{YEAR}-01-01" and not by_prefix:
                old += 1
                if old >= 3:
                    break
                continue
            old = 0
        out.append((wb, f))
    keep = {f.name for _, f in out}
    for f in TABS_DIR.glob("*.xlsx"):  # 이번에 안 쓴 예전 탭 파일 정리
        if f.name not in keep and f.stat().st_mtime < __import__("time").time() - 86400:
            f.unlink()
    return out


def download(tries=4):
    """시트 xlsx 받기 — 파일이 커서 중간에 끊기면 다시 시도"""
    import time
    import zipfile
    url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
    for i in range(1, tries + 1):
        try:
            with urllib.request.urlopen(url, timeout=600) as r:
                XLSX.write_bytes(r.read())
            zipfile.ZipFile(XLSX).testzip()  # 덜 받은 파일 걸러내기
            return
        except Exception as e:
            print(f"download try {i} failed: {e}")
            if i == tries:
                raise
            time.sleep(20 * i)


def norm(s):
    return re.sub(r"[\s/()]|\d차", "", s or "")


def text(v):
    if v is None or isinstance(v, bool):
        return ""
    if isinstance(v, dt.datetime):
        return f"{v.month}/{v.day}"
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def fill(cell):
    return cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None


def iso(d):
    return d.strftime("%Y-%m-%d")


def shown(ws, r, c, plain):
    """화면에 보일 칸 글자: 시트의 굵게·색·크기·링크 표시를 넣어서 돌려준다 (구조 판단은 plain 으로)"""
    f = FORMATS.get((ws.title.strip(), r, c))
    if not f or not plain:
        return plain
    lead = len(f.text) - len(f.text.lstrip())
    if f.text.strip() == plain:
        return f.render(lead, lead + len(plain))
    return f.render_text(plain)


def notes_of(ws, r, cells, min_len=40):
    """행의 긴 글(목표·참고사항 등)을 서식 포함해서"""
    return [shown(ws, r, c, v) for c, v in cells if len(v) > min_len]


def row_image_rows(ws, r, cols=None):
    """r 행의 사진을 시트 배치대로: 위→아래 '사진 줄', 줄 안에서는 왼쪽→오른쪽. [(파일, 줄 번호)]"""
    cells = IMAGES.get(ws.title.strip(), {})
    shots = [x for (rr, c), lst in cells.items() if rr == r and (cols is None or c in cols) for x in lst]
    if not shots:
        return []
    shots.sort(key=lambda x: x[1])
    bands, cur, top, hmin = [], [], None, None
    for s in shots:  # 앞 줄의 사진 높이 절반 이상 아래에서 시작하면 새 줄
        if cur and s[1] > top + 0.5 * hmin:
            bands.append(cur)
            cur, top, hmin = [], None, None
        cur.append(s)
        top = s[1] if top is None else top
        hmin = s[2] if hmin is None else min(hmin, s[2])
    bands.append(cur)
    return [(f"media/{s[0]}", b) for b, band in enumerate(bands) for s in sorted(band, key=lambda x: (x[3], x[4]))]


def row_shots(ws, r):
    """r 행 사진: [(파일, (행 안 세로 위치 EMU, 사진 높이 EMU))] (시트 배치 순서)"""
    cells = IMAGES.get(ws.title.strip(), {})
    by_file = {}
    for (rr, c), lst in cells.items():
        if rr == r:
            for s in lst:
                by_file.setdefault(f"media/{s[0]}", (s[1], s[2]))
    return [(f, by_file.get(f, (0, 0))) for f, _ in row_image_rows(ws, r)]


NUM_RE = re.compile(r"^#?\s*(\d{1,2})\s*(?:[~\-]\s*\d{1,2}\s*)?[.)]")
CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"


def numbered_starts(text):
    """번호가 1, 2, 3 … 차례로 시작하는 줄 번호 (화면 template 의 numberedStarts 와 같은 규칙)"""
    starts, want = [], 1
    for i, ln in enumerate(text.split("\n")):
        t = re.sub(r"⟦([^|⟧]*)\|[^⟧]*⟧", r"\1", re.sub(r"⟪[^⟫]*⟫", "", ln)).strip()
        m = NUM_RE.match(t)
        num = int(m[1]) if m else (CIRCLED.index(t[0]) + 1 if t and t[0] in CIRCLED else None)
        if num == want:
            starts.append(i)
            want += 1
    return starts


PHOTO_ANCHOR = float(os.environ.get("PHOTO_ANCHOR", "0.5"))  # 사진 가운데 높이 옆의 항목에 붙임


def est_layout(text, width_px, font_pt):
    """글 줄마다 시작 높이(pt) 추정: 칸 너비로 자동 줄바꿈 (한글 ≈ 1em, 영문·숫자 ≈ 0.55em, 공백 ≈ 0.3em)"""
    font_px, line_pt = font_pt * 4 / 3, font_pt * 1.2
    usable = max(40, width_px - 8)
    tops, y = [], 0.0
    for ln in text.split("\n"):
        tops.append(y)
        w = sum(font_px * (0.3 if ch == " " else 0.55 if ord(ch) < 128 else 1.0) for ch in ln)
        y += line_pt * max(1, -(-w // usable))
    return tops, y


def cell_width_px(ws, row, col):
    """칸 화면 너비(px) — 시트 HTML 보기 기준, 가로 병합이면 합"""
    merged = next((m for m in ws.merged_cells.ranges if m.min_row == row and m.min_col == col), None)
    widths = sheet_media.WIDTHS.get(sheet_media.sheet_key(ws.title.strip()), {})
    return sum(widths.get(c) or (ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width or 12.63) * 7.5
               for c in range(col, (merged.max_col if merged else col) + 1))


def plain_cell(v):
    return re.sub(r"⟦([^|⟧]*)\|[^⟧]*⟧", r"\1", re.sub(r"⟪[^⟫]*⟫", "", str(v)))


def place_photos(ws, item):
    """사진이 시트에서 놓인 높이로, 옆 칸 번호 항목(#1, #2 …) 중 어디에 붙는지 계산.
    결과: item["__field"] = 짝지을 칸, item["__place"] = 사진별 항목 번호(-1 = 첫 항목 앞)"""
    imgs, pos, src = item.get("__images", []), item.pop("__pos", []), item.pop("__src", {})
    item.pop("__bands", None)
    if not imgs or len(pos) != len(imgs):
        return
    cands = [(k, v) for k, v in item.items() if not k.startswith("__") and isinstance(v, str) and k in src
             and len(numbered_starts(v)) >= 2]
    if not cands:
        return
    field, text = next(((k, v) for k, v in cands if re.search(r"비주얼|가이드|설명|구성|플로우", k)), cands[0])
    row, col = src[field]
    cell = ws.cell(row, col)
    merged = next((m for m in ws.merged_cells.ranges if m.min_row == row and m.min_col == col), None)
    width_px = cell_width_px(ws, row, col)
    font_pt = float(cell.font.sz or 10)
    tops, text_h = est_layout(plain_cell(text), width_px, font_pt)
    heights = [ws.row_dimensions[r].height for r in range(row, (merged.max_row if merged else row) + 1)]
    cell_h = sum(heights) if heights and all(heights) else None
    last_row = max(r for r, _ in pos)
    if cell_h is None or any(ws.row_dimensions[x].height is None for x in range(row, last_row + 1)):
        return  # 행 높이가 자동이면 시트에 적힌 사진 위치를 믿을 수 없음 → 끝에 시트 순서대로
    if not merged:  # 행 높이는 가장 긴 글 칸에 맞춰져 있으므로, 그 칸 추정 높이로 줄 간격을 보정
        tallest = max((est_layout(plain_cell(ws.cell(row, c).value), cell_width_px(ws, row, c),
                                  float(ws.cell(row, c).font.sz or 10))[1]
                       for c in range(1, ws.max_column + 1) if isinstance(ws.cell(row, c).value, str)), default=0)
        if tallest > 0:
            scale = min(1.6, max(0.6, (cell_h - 4) / tallest))
            tops, text_h = [t * scale for t in tops], text_h * scale
    valign = (cell.alignment.vertical or "bottom") if cell_h and cell_h > text_h else "top"
    offset = 0 if valign == "top" else (cell_h - text_h) / (2 if valign == "center" else 1)
    starts = [offset + tops[i] for i in numbered_starts(text)]
    place = []
    for r, (off, h) in pos:
        if r < row:
            place.append(-1)
            continue
        above = sum((ws.row_dimensions[x].height or (text_h + offset if x == row else 15.75)) for x in range(row, r))
        y_img = above + (off + PHOTO_ANCHOR * h) / 12700
        k = max([i for i, s in enumerate(starts) if s <= y_img], default=-1)
        place.append(k)
    item["__field"], item["__place"] = field, place


def row_images(ws, r, cols=None):
    """r 행에 걸린 이미지 (시트 배치 순서)"""
    return [f for f, _ in row_image_rows(ws, r, cols)]


def top_left_only(ws):
    """merged range 의 좌상단이 아닌 셀 좌표 집합 (값 없음 취급)"""
    return {(r, c) for m in ws.merged_cells.ranges
            for r in range(m.min_row, m.max_row + 1)
            for c in range(m.min_col, m.max_col + 1)
            if (r, c) != (m.min_row, m.min_col)}


# ---------------------------------------------------------------- 캘린더 탭
def parse_calendar(ws, product_keys):
    legend, events = {}, []
    for row in ws.iter_rows(min_col=7, max_col=8):
        g, h = row
        if not text(g.value) and fill(g) and norm(text(h.value)) in product_keys:
            legend[fill(g)] = text(h.value)

    merged = {(m.min_row, m.min_col): m for m in ws.merged_cells.ranges}
    month, week_rows, header_rows, cur_year = None, [], set(), None
    block_year, prev_month, direction = YEAR, None, 0
    # B열에 "9월" 처럼 달 제목만 있는 달력은 지난 달력(옛 연도)을 행 숨기기로 가려 둔 시트 → 가려진 행은 읽지 않음
    hidden = {r for r in range(1, ws.max_row + 1) if ws.row_dimensions[r].hidden}
    if not any(r not in hidden and re.fullmatch(r"\d{1,2}월", text(ws.cell(r, 2).value)) for r in range(1, ws.max_row + 1)):
        hidden = set()
    for r in range(1, ws.max_row + 1):
        if r in hidden:
            continue
        sched = re.fullmatch(r"(\d{1,2})월(?:\s*스케줄)?", text(ws.cell(r, 2).value))
        if sched:  # "10월 스케줄", "9월 스케줄", "9월" … 처럼 월별 작은 달력이 이어진 탭
            m_ = int(sched[1])
            if prev_month is not None and m_ != prev_month:
                direction = direction or (1 if m_ > prev_month else -1)
                if direction < 0 and m_ > prev_month:
                    block_year -= 1
                if direction > 0 and m_ < prev_month:
                    block_year += 1
            month, prev_month = m_, m_
            header_rows.add(r)
            continue
        years = {v.year for v in (ws.cell(r, c).value for c in range(2, 19)) if isinstance(v, dt.datetime)}
        if years:
            cur_year = max(years)
        b = ws.cell(r, 2).value
        month_label = text(ws.cell(r, 12).value)
        if re.fullmatch(r"\d{1,2}월", month_label):
            month = int(month_label[:-1])
            header_rows.add(r)
            continue
        dates = {}
        for i, c in enumerate(CAL_COLS):
            v = ws[f"{c}{r}"].value
            if isinstance(v, dt.datetime) and v.year == YEAR:
                dates[i] = v.day
            elif (isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer() and 1 <= v <= 31
                  and cur_year in (None, YEAR)):
                dates[i] = int(v)
        if (len(dates) == 1 and not any(isinstance(ws[f"{c}{r}"].value, dt.datetime) for c in CAL_COLS)
                and not (prev_month and list(dates.values()) == [1])):
            dates = {}  # 숫자 하나만 있는 행은 날짜 행으로 보지 않음 (스케줄형 달력의 1일은 예외)
        if dates and month:
            week_rows.append((r, month, dates, block_year if prev_month else YEAR))

    for idx, (r, month, dates, year) in enumerate(week_rows):
        nxt = week_rows[idx + 1][0] if idx + 1 < len(week_rows) else ws.max_row + 1
        for rr in range(r + 1, min(nxt, r + 7)):
            if rr in header_rows:
                break
            if rr in hidden:
                continue
            for i, c in enumerate(CAL_COLS):
                cell = ws[f"{c}{rr}"]
                t = text(cell.value)
                if not t or isinstance(cell.value, dt.datetime) or i not in dates:
                    continue
                # 셀에 적힌 날짜가 월 헤더와 다른 경우(5월 블록 오류)도 월 헤더 기준으로 보정
                if not _valid(year, month, dates[i]):
                    continue
                start = dt.date(year, month, dates[i])
                span = 1
                m = merged.get((rr, cell.column))
                if m:
                    span = m.max_col - m.min_col + 1
                f = fill(cell)
                kind = "meeting" if f == MEETING_FILL else "holiday" if f == HOLIDAY_FILL else "content"
                events.append({
                    "date": iso(start),
                    "end": iso(start + dt.timedelta(days=span - 1)),
                    "text": t,
                    "product": legend.get(f),
                    "kind": kind,
                    "order": rr - r,
                    "_fill": f,
                })
    colors = {name: "#" + rgb[2:] for rgb, name in legend.items()}
    return events, colors


# ---------------------------------------------------------------- 제품 탭
DATE_RE = re.compile(r"(\d{1,2})\s*/\s*(\d{1,2})")
FORMAT_RE = re.compile(r"(스토리|릴스|게시글|게시물|캐러셀|피드|무물|라이브|라방)")


WEEKDAYS = "월화수목금토일"
TAB_YM_RE = re.compile(r"^\W*(\d{2})[._](\d{1,2})(?!\d)")  # 탭 이름 앞 "26.09", "★26.09", "26.08_" …
WD_RE = re.compile(r"(\d{1,2})\s*(?:/|월)\s*(\d{1,2})\s*일?\s*\(?\s*([월화수목금토일])")


def by_weekday(mo, d, wd):
    """적힌 요일과 맞는 해 (config 기본 연도에 가장 가까운 해) — 연도 없는 옛 탭용"""
    years = [y for y in range(YEAR + 1, YEAR - 6, -1) if _valid(y, mo, d) and dt.date(y, mo, d).weekday() == WEEKDAYS.index(wd)]
    return dt.date(min(years, key=lambda y: (abs(y - YEAR), -y)), mo, d) if years else None


def _valid(y, mo, d):
    try:
        dt.date(y, mo, d)
        return True
    except ValueError:
        return False


def dater(tab_name):
    """탭 이름이 'YY.MM…' 이면 그 연월에 가장 가까운 해로, 아니면 기본 연도로 (월, 일[, 요일]) → date"""
    m = re.match(TAB_YM_RE, tab_name)
    weekday_mode = CONFIG.get("infer_year") == "weekday"
    if not m:
        def make(mo, d, wd=None):
            if weekday_mode and wd:
                return by_weekday(mo, d, wd) or dt.date(YEAR, mo, d)
            return dt.date(YEAR, mo, d)
        return make
    ty, tm = 2000 + int(m[1]), int(m[2])

    def make(mo, d, wd=None):
        anchor = dt.date(ty, tm, 1)
        return min((dt.date(y, mo, d) for y in (ty - 1, ty, ty + 1)), key=lambda x: abs((x - anchor).days))
    return make


def parse_period(s, date_of=None):
    date_of = date_of or (lambda mo, d, wd=None: dt.date(YEAR, mo, d))
    if CONFIG.get("infer_year") == "weekday":
        w = re.search(r"(\d{1,2})\s*(?:/|월)\s*(\d{1,2})\s*일?\s*\(\s*([월화수목금토일])[^)]*\)[^~\n]{0,12}?[~\-–]\s*"
                      r"(\d{1,2})\s*(?:/|월)\s*(\d{1,2})", s)
        if w:
            a = date_of(int(w[1]), int(w[2]), w[3])
            b = dt.date(a.year + (int(w[4]) < int(w[1])), int(w[4]), int(w[5]))
            return [iso(a), iso(b)]
    m = re.search(r"(\d{1,2})/(\d{1,2})\s*\([^)]*\)\s*~\s*(\d{1,2})/(\d{1,2})", s)
    if not m:
        m = re.search(r"\d{2}\.(\d{2})\.(\d{2})\s*~\s*\d{2}\.(\d{2})\.(\d{2})", s)
    if not m:
        return None
    a = date_of(int(m[1]), int(m[2]))
    b = date_of(int(m[3]), int(m[4]))
    return [iso(a), iso(b)]


def split_topic(b):
    lines = [l.strip() for l in b.split("\n") if l.strip()]
    if not lines:
        return "", "", []
    fmt = lines[0]
    m = re.match(r"\[(.+?)\]\s*(.*)", fmt)
    if m:
        fmt = m[1].strip()
        if m[2].strip():
            lines = [fmt, m[2].strip(), *lines[1:]]
    if not FORMAT_RE.match(fmt):
        return "기타", " ".join(lines), []
    if fmt.startswith("스토리"):
        rest = " ".join(l.lstrip("*").strip() for l in lines[1:])
        return "스토리", "", [rest] if rest else []
    title = [l for l in lines[1:] if not l.startswith("*")]
    notes = [l.lstrip("*").strip() for l in lines[1:] if l.startswith("*")]
    return fmt, " ".join(title), notes


DATE_HEADERS = ("일정", "업로드 일자")


def find_header(ws):
    """(헤더 행, 날짜 열) — 제품 탭이 아니면 None"""
    for r in range(1, 16):
        for c in (1, 2):
            if text(ws.cell(r, c).value) in DATE_HEADERS:
                return r, c
    for r in range(1, 16):  # '구분 | 일정 | 노출방식 …' 처럼 한 칸 밀린 표
        if text(ws.cell(r, 3).value) == "일정" and text(ws.cell(r, 4).value) == "노출방식":
            return r, 3
    for r in range(1, 16):  # 날짜 칸 제목이 비어 있는 표: '피드 주제' 두 칸 왼쪽이 날짜
        for c in (3, 4):
            if text(ws.cell(r, c).value) == "피드 주제":
                return r, c - 2
    return None


def slot_label(a):
    """'D-12\n9/26(토)' → 'D-12', '9/29 오픈 (화)' → 'OPEN', 날짜뿐이면 ''"""
    m = re.search(r"D\s*[-+]\s*\d+", a, re.I)
    if m:
        return re.sub(r"\s+", "", m[0]).upper()
    m = re.search(r"OPEN|오픈(?!\s*[전후])|마감", a, re.I)
    if m:
        return "마감" if m[0] == "마감" else "OPEN"
    m = re.match(r"\s*오픈\s*([전후])", a)
    if m:
        return f"오픈 {m[1]}"
    first = a.split("\n")[0].split(" ")[0].strip()
    return "" if DATE_RE.match(first) else first


def kind_topic(kind, title):
    """업로드 일자형 탭: (형식, 제목, 메모)"""
    kind, title = kind.strip(), title.strip()
    m = re.match(r"\[(.+?)\]\s*(.*)", title, re.S)
    if m and not kind:
        kind, title = m[1], m[2]
    both = kind + " " + title
    if "스토리" in kind or (not kind and "스토리" in title):
        fmt = "스토리"
    elif FORMAT_RE.search(kind):
        fmt = FORMAT_RE.search(kind)[1]
    elif "일상" in both or "자율" in both:
        fmt = "자유일상"
    else:
        fmt = "기타"
    title = " ".join(l.strip() for l in title.split("\n") if l.strip())
    if not title and fmt not in ("스토리",) and kind and not FORMAT_RE.search(kind):
        title = kind
    return fmt, title, []


def parse_product(ws, colors):
    skip = top_left_only(ws)

    def val(r, c):
        return "" if (r, c) in skip else text(ws.cell(r, c).value)

    header_row, date_col = find_header(ws)
    head = {}
    for c in range(date_col, ws.max_column + 1):
        h = val(header_row, c)
        if h:
            n = sum(1 for x in head.values() if x == h or re.fullmatch(re.escape(h) + r" \d+", x))
            head[c] = h if n == 0 else f"{h} {n + 1}"
    head.setdefault(date_col, "업로드 일자")  # 날짜 칸 제목이 비어 있으면 업로드 일자형으로
    upload_style = head[date_col] == "업로드 일자"
    date_of = dater(ws.title)
    if upload_style:  # 업로드 일자 | 촬영 기한 | 콘텐츠 유형 | 콘텐츠 주제 | …
        kind_col = next((c for c, h in head.items() if "유형" in h), None)
        title_col = next((c for c, h in head.items() if "주제" in h), None)
        topic_cols = [c for c in (kind_col, title_col) if c]
        headers = {c: h for c, h in head.items() if c != date_col and c not in topic_cols}
    else:  # 일정 | 주제 | 제작 의도 | …   또는   일정 | 노출방식 | 비주얼 설명 | 비주얼
        topic_cols = [date_col + 1]
        headers = {c: h for c, h in head.items() if c > date_col + 1}
    expose_style = head.get(date_col + 1) == "노출방식"

    info, period, product_line, top_lines = [], None, None, []
    for r in range(1, header_row):
        t = val(r, 1) or val(r, 2)
        top_lines += t.split("\n")
        c0 = 1 if val(r, 1) else 2
        if not t:
            continue
        pair = [(c, val(r, c)) for c in range(1, 7) if val(r, c)]
        if expose_style and len(pair) >= 2 and len(pair[0][1]) <= 30 and "📍" not in pair[0][1]:
            label = re.sub(r"\s*\n\s*", " ", pair[0][1]).strip()
            info.append({"label": label, "text": shown(ws, r, pair[1][0], pair[1][1])})
            for c, v in pair[2:]:
                if len(v) > 40:
                    info.append({"label": "참고", "text": shown(ws, r, c, v)})
            top_lines += pair[1][1].split("\n")
            if "공구" in label and ("진행" in label or "일정" in label):
                period = period or parse_period(pair[1][1], date_of)
            if "제품" in label:
                product_line = product_line or pair[1][1].split("\n")[0].strip()
            continue
        if "📍" not in t and "\n" not in t:  # 업로드 일자형: "방효선 X 쑥세럼&크림 8/26(수) ~ 8/30(일)"
            period = period or parse_period(t, date_of)
            if len(t.strip()) > 4:
                info.append({"label": "공구", "text": shown(ws, r, c0, t)})
            continue
        for block in re.split(r"\n\s*\n(?=📍)", t):
            block = block.strip()
            head_line, _, body = block.partition("\n")
            head_line = head_line.replace("📍", "").strip()
            if "공구일정" in head_line:
                period = parse_period(head_line, date_of)
            if "공구상품" in head_line:
                product_line = head_line.split(":", 1)[-1].strip()
            if ":" in head_line and not body:
                k, v = head_line.split(":", 1)
                info.append({"label": k.strip(), "text": shown(ws, r, c0, v.strip())})
            else:
                label, _, rest = head_line.partition(":")
                body = (rest.strip() + "\n" + body).strip() if rest.strip() else body
                info.append({"label": label.strip(), "text": shown(ws, r, c0, body.strip())})

    # 날짜·주제 칸은 병합이 많아 값 전파
    merge_top = {}
    for m in ws.merged_cells.ranges:
        for r in range(m.min_row, m.max_row + 1):
            if m.min_col in (date_col, *topic_cols):
                merge_top[(r, m.min_col)] = m.min_row

    def cell_text(r, c):
        return text(ws.cell(merge_top.get((r, c), r), c).value) if c else ""

    def is_new(r, c, v):
        return bool(v) and merge_top.get((r, c), r) == r

    info_images = [i for r in range(1, header_row) for i in row_images(ws, r)]
    pending = [(f, header_row, (0, h)) for f, (_, h) in row_shots(ws, header_row)]  # 헤더 줄에 걸친 이미지는 첫 행 것
    slots, cur, last_a = [], None, ""
    for r in range(header_row + 1, ws.max_row + 1):
        a = cell_text(r, date_col)
        tops = [cell_text(r, c) for c in topic_cols]
        b = "\n\n".join(t for t in tops if t)
        fields, srcs = {}, {}
        for c, h in headers.items():
            v = val(r, c)
            if v and v != "-":
                fields[h] = shown(ws, r, c, v)
                srcs[h] = (r, c)
        shots = [(f, r, off) for f, off in row_shots(ws, r)]
        imgs = pending + shots
        pending = []
        if not (a or b or fields or imgs):
            continue
        new_a = is_new(r, date_col, a)
        new_b = any(is_new(r, c, t) for c, t in zip(topic_cols, tops))
        if cur is None or new_a or new_b:
            a = a or last_a
            last_a = a
            m = DATE_RE.search(a)
            if upload_style:
                if len(topic_cols) == 2:
                    fmt, title, notes = kind_topic(*(tops + ["", ""])[:2])
                elif topic_cols and topic_cols[0] == kind_col:
                    fmt, title, notes = kind_topic(b, "")
                    title = "" if title == b.strip() and fmt != "기타" else title  # 형식 글자(캐러셀 등)는 제목으로 쓰지 않음
                else:
                    fmt, title, notes = kind_topic("", b)
            elif expose_style:
                found = FORMAT_RE.search(b)
                fmt, title, notes = (found[1] if found else "기타"), "", []
            else:
                fmt, title, notes = split_topic(b)
            cur = {
                "label": slot_label(a),
                "date": iso(date_of(int(m[1]), int(m[2]), (WD_RE.search(a) or [None] * 4)[3])) if m else None,
                "format": fmt or "기타",
                "title": title,
                "notes": notes,
                "items": [],
            }
            if expose_style and not cur["label"]:
                cur["label"] = slot_label(b) if re.search(r"D\s*[-+]\s*\d+|OPEN|오픈|마감", b, re.I) else ""
            if upload_style and not cur["label"] and re.fullmatch(r"(D\s*[-+]\s*\d+|OPEN|오픈|마감)(\s*피드)?", title, re.I):
                cur["label"] = slot_label(title)  # 제목 칸에 적힌 D-1 / OPEN / 마감
            slots.append(cur)
        if fields:
            if cur["format"].startswith("스토리") or not cur["items"]:
                cur["items"].append({**fields, "__src": srcs})
            else:
                last = cur["items"][-1]
                for k, v in fields.items():
                    last[k] = (last[k] + "\n\n" + v) if k in last else v
                    last.setdefault("__src", {}).setdefault(k, srcs[k])
        if imgs:
            if not cur["items"]:
                cur["items"].append({})
            it = cur["items"][-1]
            it.setdefault("__images", []).extend(f for f, _, _ in imgs)
            it.setdefault("__pos", []).extend((rr, off) for _, rr, off in imgs)

    def plain_of(v):
        return re.sub(r"⟪[^⟫]*⟫|⟦([^|⟧]*)\|[^⟧]*⟧", lambda x: x.group(1) or "", v)

    for sl in slots:  # 노출방식형: 비주얼 설명의 "✅주제 : …" 를 제목으로
        if not expose_style or sl["title"]:
            continue
        desc = next((plain_of(v) for it in sl["items"] for k, v in it.items() if "설명" in k), "")
        m = re.search(r"주제[^:\n]{0,12}:\s*(.+)", desc)
        line = (m[1] if m else next((ln for ln in desc.split("\n") if ln.strip()), "")).strip(" ✅●/")
        sl["title"] = line[:60]
    if expose_style and period:  # 날짜 오타(요일 불일치 등)로 연도가 튀면 공구 기간에 가까운 해로
        open_day = dt.date.fromisoformat(period[0])
        for sl in slots:
            if sl["date"]:
                d0 = dt.date.fromisoformat(sl["date"])
                cands = [dt.date(y, d0.month, d0.day) for y in (open_day.year - 1, open_day.year, open_day.year + 1)
                         if _valid(y, d0.month, d0.day)]
                sl["date"] = iso(min(cands, key=lambda x: abs((x - open_day).days)))
    for sl in slots:
        if not upload_style:
            break
        fields = [(k, plain_of(v).strip()) for it in sl["items"] for k, v in it.items() if not k.startswith("__")]
        if sl["format"] == "기타":  # 형식 칸이 없으면 내용 앞 표시로 (★피드, [스토리] …)
            head = next((v for k, v in fields if "가이드" in k or "내용" in k), "")
            if re.match(r"[★\[]\s*스토리", head) or sl["title"] in ("[스토리]", "스토리"):
                sl["format"] = "스토리"
                if sl["title"] in ("[스토리]", "스토리"):
                    sl["title"] = ""
            elif re.match(r"[★\[]\s*(자유\s*)?일상", head):
                sl["format"] = "자유일상"
            elif re.match(r"[★\[]\s*피드", head) or re.search(r"(^|\s)피드$", sl["title"]):
                sl["format"] = "피드"
        if re.fullmatch(r"(D\s*[-+]\s*\d+|OPEN|오픈|마감)(\s*피드)?", sl["title"], re.I):
            # 제목 칸에 D-1 / OPEN / 마감만 있으면 피드글 첫 줄을 제목으로
            cap = next((v for k, v in fields if "피드글 최종" in k), "") or next((v for k, v in fields if "피드글" in k), "")
            line = next((ln.strip() for ln in cap.split("\n") if ln.strip() and not ln.strip().startswith("*")), "")
            if line:
                sl["title"] = line[:40]
    for sl in slots:  # 제목이 비어 있으면 내용 첫 줄로 (예: "✨방학!!!!✨"), 주소·날짜 칸은 제외
        if not sl["title"] and sl["format"] not in ("스토리",) and upload_style:
            lines = [ln.strip() for it in sl["items"] for k, v in it.items() if not k.startswith("__") and "기한" not in k
                     for ln in re.sub(r"⟪[^⟫]*⟫|⟦([^|⟧]*)\|[^⟧]*⟧", lambda x: x.group(1) or "", v).split("\n")]
            named = next((m[2].strip() for ln in lines for m in [re.search(r"(타이틀|주제)\s*:\s*(.+)", ln)] if m), "")
            caption = [ln.strip() for it in sl["items"] for k, v in it.items() if "캡션" in k or "피드글" in k
                       for ln in plain_of(v).split("\n") if ln.strip() and not ln.strip().startswith("*")]
            lines = ([named] if named else []) + caption[:1] + lines
            lines += [ln.strip() for it in sl["items"] for k, v in it.items() if "기한" in k
                      for ln in re.sub(r"⟪[^⟫]*⟫|⟦([^|⟧]*)\|[^⟧]*⟧", lambda x: x.group(1) or "", v).split("\n")]
            line = next((ln for ln in lines if ln and not ln.startswith("http") and not DATE_RE.fullmatch(ln)), "")
            line = re.sub(r"^\d{1,2}[.)]\s*", "", line)  # "1. 효소 들고 있는 …" → 번호 떼기
            sl["title"] = line[:40] or ("참고 링크" if any(ln.startswith("http") for ln in lines) else "")
    if period is None:
        period = next((pp for pp in (parse_period(t, date_of) for t in top_lines) if pp), None)
    if period:
        open_day = dt.date.fromisoformat(period[0])
        for sl in slots:
            if sl["date"]:
                continue
            off = 0 if sl["label"] == "OPEN" else int(sl["label"][1:]) if re.fullmatch(r"D[-+]\d+", sl["label"]) else None
            if off is not None:
                sl["date"] = iso(open_day + dt.timedelta(days=off))
    if period is None:  # 기간 표기가 없으면 OPEN ~ 마지막 D+ 날짜
        opens = [sl["date"] for sl in slots if sl["label"] == "OPEN" and sl["date"]]
        if opens:
            after = [sl["date"] for sl in slots if (sl["label"].startswith("D+") or sl["label"] == "마감") and sl["date"]]
            period = [opens[0], max(after + opens)]

    for sl in slots:
        for it in sl["items"]:
            place_photos(ws, it)
            it.pop("__src", None)
            it.pop("__pos", None)
    name = ws.title.strip()
    base = norm(name)
    color = next((c for n, c in colors.items() if norm(n) == base), None)
    if color is None:
        color = PALETTE[sum(map(ord, base)) % len(PALETTE)]
    round_m = re.search(r"(\d+)차", name)
    seller = CONFIG["seller"].removesuffix("님")
    display = re.sub(r"\s*\([\d.~\-\s]*\)\s*$", "", name)          # "헤베스템 13차 (929-104)" → "헤베스템 13차"
    display = re.sub(r"^\W*\d{2}[._]\d{1,2}[\s_]*", "", display)     # "26.09 방탄커피", "★26.09 효소" → 제품명
    display = re.sub(rf"^{re.escape(seller)}\s*[xX×]\s*", "", display)  # "방효선x헤어 2종" → "헤어 2종"
    return {
        "id": "p" + re.sub(r"\W", "", base) + (round_m[1] if round_m else ""),
        "name": re.sub(r"\s*\(?\d+차\)?", "", display).strip(),
        "round": f"{round_m[1]}차" if round_m else "",
        "fullName": product_line,
        "color": color,
        "period": period,
        "info": info,
        "images": info_images,
        "slots": slots,
    }


# ---------------------------------------------------------------- 월별 탭
GRID_LABELS = {"스토리", "주제", "기획 의도", "팔로워 반응", "콘텐츠 구성", "콘텐츠 플로우", "캡션 참고",
               "콘텐츠", "기대 효과", "기획 의도 & 기대 효과", "팔로워 예상 반응", "피드 비주얼", "피드 참고"}
DAY_RE = re.compile(r"^(\d{1,2})(?:\.0)?(?:\s*\((.+)\))?$")


def parse_month_grid(ws, month):
    """3~5월: 주 단위 그리드 (날짜 행 + 라벨 행)"""
    skip = top_left_only(ws)

    def val(r, c):
        return "" if (r, c) in skip else text(ws.cell(r, c).value)

    days, notes, extra = {}, [], []
    week, first_week, in_extra = None, None, False
    for r in range(1, ws.max_row + 1):
        cells = {c: val(r, c) for c in range(2, 9)}
        a = val(r, 1)
        day_hits = {c: DAY_RE.match(v) for c, v in cells.items() if v}
        day_hits = {c: m for c, m in day_hits.items() if m}
        if len(day_hits) >= 1 and (a == "날짜" or len(day_hits) >= 3 or week is None and len(day_hits) >= 1):
            week = {}
            for c, m in day_hits.items():
                d = int(m[1])
                key = iso(dt.date(YEAR, month, d))
                week[c] = key
                days.setdefault(key, {"holiday": m[2], "fields": []})
            if first_week is None:
                first_week = r
            for c, v in cells.items():
                if len(v) > 40 and c not in day_hits:
                    notes.append(shown(ws, r, c, v))
            continue
        if week is None:
            notes += notes_of(ws, r, [(1, a), *cells.items()])
            continue
        label = a.replace("\n", " ")
        if label == "날짜":
            continue
        for c, key in week.items():
            imgs = row_images(ws, r, {c})
            if imgs:
                days[key].setdefault("images", []).extend(imgs)
        if label:
            in_extra = label not in GRID_LABELS
        for c, v in cells.items():
            if len(v) < 2:
                continue
            if in_extra:
                if label:
                    extra.append({"label": label, "text": shown(ws, r, c, v)})
                    label = ""
                else:
                    extra[-1]["text"] += "\n\n" + shown(ws, r, c, v)
            elif c in week and label:
                days[week[c]]["fields"].append({"label": label, "text": shown(ws, r, c, v)})
            elif len(v) > 40:
                notes.append(shown(ws, r, c, v))
    day_list = [{"date": k, **v} for k, v in sorted(days.items()) if v["fields"] or v["holiday"] or v.get("images")]
    return {"kind": "grid", "notes": notes, "days": day_list, "extra": extra}


def parse_month_sections(ws):
    """6~8월: 카테고리 섹션별 기획 리스트"""
    skip = top_left_only(ws)

    def val(r, c):
        return "" if (r, c) in skip else text(ws.cell(r, c).value)

    notes, sections, cur, item = [], [], None, None
    cols = {2: "기획 의도", 3: "콘텐츠 주제", 4: "콘텐츠 플로우", 5: "캡션 참고", 6: "비고"}
    for r in range(2, ws.max_row + 1):
        raw_a = ws.cell(r, 1).value if (r, 1) not in skip else None
        a = text(raw_a)
        rest = {c: val(r, c) for c in range(2, 7)}
        if a.startswith("포인트"):
            continue
        if rest.get(2) == "팔로워 반응" and not a:
            continue
        imgs = row_images(ws, r)
        if item and imgs and not isinstance(raw_a, bool) and a not in ("True", "False"):
            item.setdefault("images", []).extend(imgs)
        if isinstance(raw_a, bool) or a in ("True", "False"):
            item = {"done": a == "True" or raw_a is True, "fields": {}}
            if imgs:
                item["images"] = imgs
            for c, v in rest.items():
                if v:
                    item["fields"][cols[c]] = v if c == 3 else shown(ws, r, c, v)
            item["title"] = item["fields"].pop("콘텐츠 주제", "")
            if "[" not in item["title"][:12]:  # 대괄호 포맷이 없는 건 스토리 아이디어
                if item["title"]:
                    item["fields"] = {"스토리 내용": item["title"], **item["fields"]}
                head = item["title"].split(" - ")[0] if " - " in item["title"][:20] else ""
                item["title"] = head or "스토리 아이디어"
                item["story"] = True
            if cur is None:
                cur = {"title": "", "items": []}
                sections.append(cur)
            cur["items"].append(item)
            continue
        if a and not any(rest.values()) and r > 3:
            title, _, sub = a.partition("*")
            cur = {"title": title.strip(), "sub": sub.strip(), "items": []}
            sections.append(cur)
            item = None
            continue
        if not sections:
            notes += notes_of(ws, r, [(1, a), *rest.items()])
            continue
        if item and not a:
            if rest.get(2):
                item["reaction"] = (item.get("reaction", "") + "\n" + shown(ws, r, 2, rest[2])).strip()
            for c in range(3, 7):
                if rest.get(c):
                    k = cols[c]
                    item["fields"][k] = (item["fields"].get(k, "") + "\n\n" + shown(ws, r, c, rest[c])).strip()
    return {"kind": "sections", "notes": notes, "sections": sections}


def parse_month_weeks(ws):
    """주차별 섹션: 'n주차 콘텐츠' 제목 → '유형' 헤더 → 항목 행 + 이어지는 보조 행"""
    skip = top_left_only(ws)

    def val(r, c):
        return "" if (r, c) in skip else text(ws.cell(r, c).value)

    notes, sections, headers, item = [], [], {}, None
    for r in range(1, ws.max_row + 1):
        a = val(r, 1)
        row = {c: val(r, c) for c in range(2, ws.max_column + 1)}
        if "주차" in a:
            title, _, rng = a.partition("(")
            sections.append({"title": title.replace("콘텐츠", "").strip(),
                             "sub": rng.rstrip(")").strip(), "items": []})
            item = None
            continue
        if a == "유형":
            headers = {c: v for c, v in row.items() if v}
            continue
        if not sections:
            notes += notes_of(ws, r, [(1, a), *row.items()])
            continue
        done = any(ws.cell(r, c).value is True for c in range(2, ws.max_column + 1))
        imgs = row_images(ws, r)
        if item and imgs and not a:
            item.setdefault("images", []).extend(imgs)
        if a:
            fields = {headers[c]: (v if c == 2 else shown(ws, r, c, v)) for c, v in row.items() if v and c in headers}
            topic = fields.pop(headers.get(2, ""), "")
            item = {"done": done, "title": f"[{a}] {topic}".strip(), "fields": fields}
            if imgs:
                item["images"] = imgs
            sections[-1]["items"].append(item)
        elif item:
            for c, v in row.items():
                if not v or c not in headers:
                    continue
                if "팔로워" in headers[c] and headers[c] == headers.get(3):
                    item["reaction"] = (item.get("reaction", "") + "\n" + shown(ws, r, c, v)).strip()
                else:
                    k = headers[c]
                    item["fields"][k] = (item["fields"].get(k, "") + "\n\n" + shown(ws, r, c, v)).strip()
    return {"kind": "sections", "notes": notes, "sections": sections}


def parse_month(ws):
    m = re.match(r"(\d{2})\.(\d{2})", ws.title.strip())
    month = int(m[2])
    has_grid = any(text(ws.cell(r, 1).value) == "날짜" for r in range(1, 10))
    has_weeks = any(text(ws.cell(r, 1).value) == "유형" for r in range(1, ws.max_row + 1))
    data = (parse_month_grid(ws, month) if has_grid
            else parse_month_weeks(ws) if has_weeks else parse_month_sections(ws))
    data.update({"id": f"m{month:02d}", "month": month, "name": f"{month}월"})
    return data


# ---------------------------------------------------------------- main
PAY_RE = re.compile(r"결제\s*창[^\n]*?(\d{1,2})\s*/\s*(\d{1,2})[^\n\d]*?(오전|오후)?\s*(\d{1,2})\s*(?::\s*(\d{2})|시(?:\s*(\d{1,2})\s*분)?)")


def pay_close(p):
    """제품 정보에 적힌 '결제창 마감: 5/11(월) 오전 11시 20분' → '2026-05-11T11:20' (없으면 None)"""
    if not p["period"]:
        return None
    end = dt.date.fromisoformat(p["period"][1])
    for x in p["info"]:
        raw = re.sub(r"⟦([^|⟧]*)\|[^⟧]*⟧", r"\1", re.sub(r"⟪[^⟫]*⟫", "", f'{x["label"]}\n{x["text"]}'))
        for ln in raw.split("\n"):
            m = PAY_RE.search(ln) if "마감" in ln else None
            if not m:
                continue
            mo, d, ampm, h, mi = int(m[1]), int(m[2]), m[3], int(m[4]), int(m[5] or m[6] or 0)
            if ampm == "오후" and h < 12:
                h += 12
            if ampm == "오전" and h == 12:
                h = 0
            for y in (end.year, end.year + 1, end.year - 1):  # 연말·연초에 걸친 공구
                if _valid(y, mo, d) and h < 24 and mi < 60 and -1 <= (dt.date(y, mo, d) - end).days <= 14:
                    return f"{y}-{mo:02d}-{d:02d}T{h:02d}:{mi:02d}"
    return None


def main():
    local = "--local" in sys.argv
    if CONFIG.get("per_tab"):
        books = load_tabs(local)
    else:
        if not local:
            download()
        books = [(openpyxl.load_workbook(XLSX, data_only=True, rich_text=True), XLSX)]
    try:
        links = sheet_media.fetch_links(SHEET_ID)
    except Exception as e:  # 시트 HTML 보기를 못 읽으면 xlsx 의 셀 링크만 사용
        print("links: htmlview 실패, xlsx 링크로 대체 -", e)
        links = {}
    sheets = []
    for wb, path in books:
        FORMATS.update(sheet_media.collect_formats(wb, links))
        visible = {ws.title.strip() for ws in wb.worksheets if ws.sheet_state == "visible"}
        IMAGES.update(sheet_media.extract_images(path, MEDIA, only=visible))
        sheets += [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
    print("formats:", {k: v for k, v in sheet_media.STATS.items()})
    print("tabs:", len(sheets), "images:", sum(len(v) for cells in IMAGES.values() for v in cells.values()))
    cal = next((ws for ws in sheets if "캘린더" in ws.title or "스케줄" in ws.title), None)
    product_sheets = [ws for ws in sheets if ws is not cal and find_header(ws)]
    month_sheets = [ws for ws in sheets if ws not in product_sheets and re.match(r"\d{2}\.\d{2}", ws.title.strip())]
    events, colors = parse_calendar(cal, {norm(ws.title) for ws in product_sheets}) if cal else ([], {})
    products = [parse_product(ws, colors) for ws in product_sheets]
    for p, ws in zip(products, product_sheets):  # 이름이 같은 제품(차수 표기 없음)은 탭의 연월로 구분
        tab_ym = TAB_YM_RE.match(ws.title)
        if not p["round"] and tab_ym and sum(q["name"] == p["name"] for q in products) > 1:
            p["round"] = f"{tab_ym[1]}.{int(tab_ym[2]):02d}"
    months = [parse_month(ws) for ws in month_sheets]
    for e in events:  # 범례 색이 없으면 "제품명 + N차" 가 적힌 칸을 그 제품으로
        if not e["product"]:
            hit = next((p for p in products if p["round"] and norm(p["name"]) in norm(e["text"])
                        and p["round"] in e["text"].replace(" ", "")), None)
            if hit:
                e["product"], e["_id"] = hit["name"], hit["id"]
    def near(p, e):  # 제품 탭의 날짜들과 캘린더 칸 사이 거리(일)
        ds = [s["date"] for s in p["slots"] if s["date"]] + (p["period"] or [])
        return min((abs((dt.date.fromisoformat(d) - dt.date.fromisoformat(x)).days)
                    for d in ds for x in (e["date"], e["end"])), default=999)
    for e in events:  # 여러 날에 걸쳐 제품명만 적힌 칸("방탄커피", "아르히 헤어 2종")은 날짜가 가까운 그 제품 탭으로
        if e["product"] or e["end"] <= e["date"]:
            continue
        t = norm(e["text"])
        cands = [p for p in products if len(norm(p["name"])) >= 2 and (norm(p["name"]) in t or t in norm(p["name"]))
                 and near(p, e) <= 21]
        if cands:
            hit = min(cands, key=lambda p: near(p, e))
            e["product"], e["_id"] = hit["name"], hit["id"]
    for e in events:  # 범례 이름 → 제품 탭 id (같은 제품 여러 차수면 날짜가 가까운 차수)
        cands = [p for p in products if e["product"] and norm(p["name"]) == norm(e["product"])]

        def dist(p):
            ds = [s["date"] for s in p["slots"] if s["date"]] + (p["period"] or [])
            return min(abs((dt.date.fromisoformat(d) - dt.date.fromisoformat(e["date"])).days) for d in ds)
        e["product"] = e.pop("_id", None) or (min(cands, key=dist)["id"] if cands else None)
    # 캘린더에 여러 날에 걸쳐 제품명이 적힌 칸 = 그 기간에 공구. 제품 탭이 아직 없어도 같은 색 칸이면 공구로 봄
    by_id = {p["id"]: p for p in products}
    for p in products:  # 결제창 마감 시각이 따로 적힌 공구
        pc = pay_close(p)
        if pc:
            p["payClose"] = pc
    sale_fills = {e["_fill"] for e in events if e["end"] > e["date"] and e["product"]
                  and e["_fill"] not in (None, MEETING_FILL, HOLIDAY_FILL)}
    for e in events:
        f = e.pop("_fill", None)
        if e["end"] > e["date"] and (e["product"] or f in sale_fills):
            e["sale"] = True
            p = by_id.get(e["product"])
            if p and not p["period"]:  # 제품 탭에 공구일정이 안 적혀 있으면 캘린더의 기간을 사용
                p["period"] = [e["date"], e["end"]]
    products.sort(key=lambda p: (p["period"] or [max((s["date"] for s in p["slots"] if s["date"]), default="0000")])[0],
                  reverse=True)
    months.sort(key=lambda m: m["month"], reverse=True)
    data = {
        "updated": dt.datetime.now(dt.timezone(dt.timedelta(hours=9))).strftime("%Y-%m-%d %H:%M"),
        "sheetUrl": f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit",
        "year": YEAR,
        "events": events,
        "products": products,
        "months": months,
    }
    (HERE / "data.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    if TEMPLATE.exists():
        html = TEMPLATE.read_text(encoding="utf-8")
        payload = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
        html = html.replace("__APP_NAME__", CONFIG["app_name"]).replace("__SELLER__", CONFIG["seller"])
        OUT.write_text(html.replace("/*__DATA__*/null", payload), encoding="utf-8")
        print("built", OUT)
    print(f"events={len(events)} products={len(products)} months={len(months)}")


if __name__ == "__main__":
    main()
